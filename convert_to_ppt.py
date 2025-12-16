"""
마크다운 파일과 Mermaid 다이어그램을 PowerPoint 프레젠테이션으로 변환하는 스크립트
"""
import os
import re
import subprocess
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

def check_dependencies():
    """필요한 도구가 설치되어 있는지 확인"""
    try:
        import pptx
        print("✓ python-pptx 설치 확인됨")
    except ImportError:
        print("✗ python-pptx 설치 필요")
        print("  설치: pip install python-pptx")
        return False
    
    try:
        subprocess.run(['mmdc', '--version'], capture_output=True, check=True)
        print("✓ Mermaid CLI 설치 확인됨")
    except (subprocess.CalledProcessError, FileNotFoundError):
        print("✗ Mermaid CLI 설치 필요")
        print("  설치: npm install -g @mermaid-js/mermaid-cli")
        return False
    
    return True

def convert_mermaid_to_image(diagram_code, output_path):
    """Mermaid 다이어그램을 이미지로 변환"""
    # 임시 mermaid 파일 생성
    temp_mmd = 'temp_diagram.mmd'
    with open(temp_mmd, 'w', encoding='utf-8') as f:
        f.write(diagram_code)
    
    # 이미지로 변환
    cmd = f'mmdc -i {temp_mmd} -o {output_path} -b transparent -w 1920 -H 1080'
    
    try:
        subprocess.run(cmd.split(), check=True, capture_output=True)
        return True
    except subprocess.CalledProcessError:
        return False
    finally:
        if os.path.exists(temp_mmd):
            os.remove(temp_mmd)

def clean_html_tags(text):
    """HTML 태그 제거"""
    # HTML 태그 제거
    text = re.sub(r'<[^>]+>', '', text)
    # HTML 엔티티 디코딩
    text = text.replace('&nbsp;', ' ')
    text = text.replace('&lt;', '<')
    text = text.replace('&gt;', '>')
    text = text.replace('&amp;', '&')
    return text.strip()

def parse_markdown(md_file, image_dir=None):
    """마크다운 파일 파싱"""
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # YAML front matter 제거
    if content.startswith('---'):
        parts = content.split('---', 2)
        if len(parts) >= 3:
            content = parts[2].strip()
    
    # HTML 태그 제거
    content = re.sub(r'<div[^>]*>', '', content)
    content = re.sub(r'</div>', '', content)
    content = re.sub(r'<style[^>]*>.*?</style>', '', content, flags=re.DOTALL)
    
    # Mermaid 다이어그램을 이미지로 변환
    if image_dir:
        pattern = r'```mermaid\s*\n(.*?)```'
        matches = list(re.finditer(pattern, content, re.DOTALL))
        
        for i, match in enumerate(matches):
            diagram_code = match.group(1)
            
            # 임시 mermaid 파일 생성
            temp_mmd = os.path.join(image_dir, f'temp_{i}.mmd')
            with open(temp_mmd, 'w', encoding='utf-8') as f:
                f.write(diagram_code)
            
            # 이미지로 변환
            output_image = os.path.join(image_dir, f'diagram_{i}.png')
            try:
                subprocess.run(
                    ['mmdc', '-i', temp_mmd, '-o', output_image, '-b', 'transparent', '-w', '1920', '-H', '1080'],
                    capture_output=True,
                    check=True
                )
                if os.path.exists(output_image):
                    rel_path = os.path.relpath(output_image, os.path.dirname(md_file))
                    image_markdown = f"![Mermaid Diagram]({rel_path})"
                    content = content[:match.start()] + image_markdown + content[match.end():]
                    print(f"  ✓ 다이어그램 {i+1} 변환 완료")
            except subprocess.CalledProcessError:
                print(f"  ✗ 다이어그램 {i+1} 변환 실패")
            finally:
                if os.path.exists(temp_mmd):
                    os.remove(temp_mmd)
    
    slides = []
    current_slide = {'title': '', 'content': [], 'level': 0}
    
    lines = content.split('\n')
    for line in lines:
        # 빈 줄 또는 구분선 무시
        if not line.strip() or line.strip() == '---':
            continue
            
        # 제목 감지
        if line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            title = line.lstrip('#').strip()
            title = clean_html_tags(title)
            
            # 이전 슬라이드 저장
            if current_slide['title'] or current_slide['content']:
                slides.append(current_slide)
            
            # 새 슬라이드 시작
            current_slide = {
                'title': title,
                'content': [],
                'level': level
            }
        # Mermaid 다이어그램 감지
        elif line.strip().startswith('```mermaid'):
            diagram_lines = []
            i = lines.index(line) + 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                diagram_lines.append(lines[i])
                i += 1
            current_slide['content'].append({
                'type': 'mermaid',
                'code': '\n'.join(diagram_lines)
            })
        # 코드 블록 제외
        elif line.strip().startswith('```'):
            continue
        # 일반 내용
        elif line.strip():
            text = clean_html_tags(line)
            if text:  # 빈 텍스트가 아닌 경우만 추가
                current_slide['content'].append({
                    'type': 'text',
                    'text': text
                })
    
    # 마지막 슬라이드 저장
    if current_slide['title'] or current_slide['content']:
        slides.append(current_slide)
    
    return slides

def get_image_size(img_path, max_width=Inches(9), max_height=Inches(5.5)):
    """이미지 크기를 슬라이드에 맞게 조정"""
    try:
        from PIL import Image
        img = Image.open(img_path)
        img_width, img_height = img.size
        
        # 비율 유지하며 크기 조정 (96 DPI 가정)
        width_ratio = max_width / Inches(img_width / 96)
        height_ratio = max_height / Inches(img_height / 96)
        ratio = min(width_ratio, height_ratio, 1.0)  # 확대하지 않음
        
        new_width = Inches(img_width * ratio / 96)
        new_height = Inches(img_height * ratio / 96)
        
        return new_width, new_height
    except (ImportError, Exception):
        # PIL이 없거나 오류 발생 시 기본 크기 사용
        return max_width, max_height

def create_presentation(slides, output_file, output_dir):
    """PowerPoint 프레젠테이션 생성 (개선된 버전)"""
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    # 색상 테마 정의
    theme_colors = {
        'primary': RGBColor(0, 51, 102),      # 진한 파랑
        'secondary': RGBColor(0, 102, 204),   # 밝은 파랑
        'accent': RGBColor(255, 102, 0),      # 주황
        'text': RGBColor(51, 51, 51),         # 진한 회색
        'light_text': RGBColor(102, 102, 102) # 밝은 회색
    }
    
    for i, slide_data in enumerate(slides):
        # 빈 슬라이드 사용 (더 많은 제어 가능)
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)
        
        # 제목 설정
        if slide_data['title']:
            title_left = Inches(0.5)
            title_top = Inches(0.3)
            title_width = Inches(9)
            title_height = Inches(0.8)
            
            title_box = slide.shapes.add_textbox(title_left, title_top, title_width, title_height)
            title_frame = title_box.text_frame
            title_frame.word_wrap = True
            title_frame.vertical_anchor = MSO_ANCHOR.TOP
            
            p = title_frame.paragraphs[0]
            p.text = slide_data['title']
            p.font.size = Pt(28 if slide_data['level'] == 1 else 24)
            p.font.bold = True
            p.font.color.rgb = theme_colors['primary']
            p.alignment = PP_ALIGN.LEFT
        
        # 내용 추가
        if slide_data['content']:
            content_top = Inches(1.3)
            content_left = Inches(0.5)
            content_width = Inches(9)
            content_height = Inches(5.8)
            
            # 이미지와 텍스트 분리
            images = [item for item in slide_data['content'] if item.get('type') == 'image']
            texts = [item for item in slide_data['content'] if item.get('type') == 'text']
            
            has_image = len(images) > 0
            has_text = len(texts) > 0
            
            # 텍스트 추가
            if has_text:
                if has_image:
                    # 이미지와 텍스트가 모두 있는 경우: 반반 레이아웃
                    text_width = Inches(4.5)
                else:
                    # 텍스트만 있는 경우
                    text_width = content_width
                
                text_box = slide.shapes.add_textbox(content_left, content_top, text_width, content_height)
                text_frame = text_box.text_frame
                text_frame.word_wrap = True
                text_frame.vertical_anchor = MSO_ANCHOR.TOP
                text_frame.margin_left = Inches(0.1)
                text_frame.margin_right = Inches(0.1)
                text_frame.margin_top = Inches(0.1)
                text_frame.margin_bottom = Inches(0.1)
                
                for content_item in texts:
                    p = text_frame.add_paragraph()
                    text = content_item['text']
                    
                    # 리스트 감지 및 처리
                    if re.match(r'^[\s]*[-*+]\s+', text) or re.match(r'^[\s]*\d+\.\s+', text):
                        text = re.sub(r'^[\s]*[-*+]\s+', '• ', text)
                        text = re.sub(r'^[\s]*\d+\.\s+', '', text)
                        p.level = 0
                        p.space_before = Pt(4)
                        p.font.size = Pt(13)
                    else:
                        p.font.size = Pt(14)
                    
                    p.text = text
                    p.space_after = Pt(8)
                    p.font.color.rgb = theme_colors['text']
                    p.alignment = PP_ALIGN.LEFT
            
            # 이미지 추가
            for img_item in images:
                img_path = img_item['path']
                if os.path.exists(img_path):
                    try:
                        # 이미지 크기 자동 조정
                        if has_text:
                            max_width = Inches(4.5)
                            max_height = Inches(5.5)
                            img_left = Inches(5)
                        else:
                            max_width = Inches(9)
                            max_height = Inches(5.8)
                            img_left = content_left
                        
                        img_width, img_height = get_image_size(img_path, max_width, max_height)
                        
                        # 이미지 위치 계산 (중앙 정렬)
                        img_top = content_top + (max_height - img_height) / 2
                        
                        slide.shapes.add_picture(
                            img_path, 
                            img_left, 
                            img_top, 
                            width=img_width,
                            height=img_height
                        )
                    except Exception as e:
                        print(f"  ⚠ 이미지 추가 실패: {img_path} - {e}")
                else:
                    print(f"  ⚠ 이미지 파일 없음: {img_path}")
        
        # 슬라이드 번호 추가
        slide_num_left = Inches(9.2)
        slide_num_top = Inches(7.2)
        slide_num_box = slide.shapes.add_textbox(slide_num_left, slide_num_top, Inches(0.8), Inches(0.3))
        slide_num_frame = slide_num_box.text_frame
        slide_num_frame.text = f"{i+1}"
        slide_num_frame.paragraphs[0].font.size = Pt(10)
        slide_num_frame.paragraphs[0].font.color.rgb = theme_colors['light_text']
        slide_num_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    
    # 프레젠테이션 저장
    prs.save(output_file)
    print(f"✓ PowerPoint 생성 완료: {output_file}")

def main():
    """메인 함수"""
    import sys
    import argparse
    from datetime import datetime
    from glob import glob
    
    # 명령줄 인자 파싱
    parser = argparse.ArgumentParser(description='마크다운 파일을 PowerPoint로 변환')
    parser.add_argument('-in', '--input-dir', type=str, default='doc',
                        help='입력 폴더 (기본값: doc)')
    parser.add_argument('-o', '--output-dir', type=str, default=None,
                        help='출력 폴더 (기본값: output_YYYYMMDD_HH-MM)')
    parser.add_argument('output', nargs='?', type=str, default=None,
                        help='출력 폴더 (하위 호환성을 위한 위치 인자)')
    
    args = parser.parse_args()
    
    # 출력 폴더 설정
    if args.output_dir:
        output_dir = args.output_dir
    elif args.output:
        output_dir = args.output
    else:
        # 현재 시각으로 폴더명 생성 (YYYYMMDD_HH-mm)
        output_dir = f"output_{datetime.now().strftime('%Y%m%d_%H-%M')}"
    
    # 입력 폴더 설정
    input_dir = args.input_dir
    
    # 출력 폴더 생성
    os.makedirs(output_dir, exist_ok=True)
    
    print("=" * 50)
    print("마크다운 → PowerPoint 변환 도구")
    if os.path.isdir(input_dir):
        print(f"입력 폴더: {input_dir}")
    print(f"출력 폴더: {output_dir}")
    print("=" * 50)
    
    # 의존성 확인
    if not check_dependencies():
        return
    
    # 변환할 마크다운 파일 목록
    if os.path.isdir(input_dir):
        # 입력 폴더에서 모든 .md 파일 찾기
        md_files = sorted(glob(os.path.join(input_dir, '**', '*.md'), recursive=True))
        if not md_files:
            print(f"✗ 입력 폴더에 .md 파일이 없습니다: {input_dir}")
            return
        print(f"  → {len(md_files)}개의 마크다운 파일을 찾았습니다")
    else:
        # 기본 파일 목록 (하위 호환성)
        md_files = [
            'doc/서비스_기획서.md',
            'doc/화면_기능_정의서.md',
            'doc/화면_흐름도.md',
        ]
    
    for md_file in md_files:
        if not os.path.exists(md_file):
            print(f"✗ 파일을 찾을 수 없습니다: {md_file}")
            continue
        
        # 이미지 저장 디렉토리 (출력 폴더 내부)
        image_dir_path = os.path.join(output_dir, 'images')
        os.makedirs(image_dir_path, exist_ok=True)
        
        print(f"\n파일 파싱 중: {md_file}")
        slides = parse_markdown(md_file, image_dir_path)
        print(f"✓ {len(slides)}개 슬라이드 생성 예정")
        
        print("\nPowerPoint 생성 중...")
        md_filename = os.path.basename(md_file).replace('.md', '.pptx')
        output_file = os.path.join(output_dir, md_filename)
        create_presentation(slides, output_file, output_dir)
        
        if os.path.exists(output_file):
            print(f"✓ 완료: {output_file}")
        else:
            print(f"✗ 실패: {md_file}")
    
    print("\n" + "=" * 50)
    print("변환 완료!")
    print(f"출력 폴더: {output_dir}")
    print("=" * 50)

if __name__ == '__main__':
    main()

